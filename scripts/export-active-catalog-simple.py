from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Uso: export-active-catalog-simple.py <origen.xlsx> <destino.xlsx>")

    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    source_book = load_workbook(source, read_only=True, data_only=True)
    source_sheet = source_book["Catalogo_utilizado"]
    headers = {str(cell.value): index for index, cell in enumerate(source_sheet[1])}

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Catalogo"
    worksheet.append(["Clave", "Concepto", "Unidad", "Precio Unitario"])

    for row in source_sheet.iter_rows(min_row=2, values_only=True):
        technical = row[headers["descripcion_tecnica"]]
        official = row[headers["descripcion_oficial"]]
        worksheet.append(
            [
                row[headers["concepto_codigo"]],
                technical or official or "",
                row[headers["unidad_simbolo"]] or row[headers["unidad_codigo"]] or "",
                float(row[headers["precio_unitario"]] or 0),
            ]
        )
    source_book.close()

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 85
    worksheet.column_dimensions["C"].width = 15
    worksheet.column_dimensions["D"].width = 20
    for row in worksheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[3].number_format = '$#,##0.00'

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    verified = load_workbook(output, read_only=True, data_only=True)
    count = verified["Catalogo"].max_row - 1
    columns = verified["Catalogo"].max_column
    verified.close()
    if count != 100 or columns != 4:
        raise RuntimeError(f"Exportación incompleta: {count} conceptos y {columns} columnas")
    print(f"OUTPUT={output}")
    print(f"ROWS={count}")
    print(f"COLUMNS={columns}")


if __name__ == "__main__":
    main()
