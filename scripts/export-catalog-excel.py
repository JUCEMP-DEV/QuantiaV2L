from __future__ import annotations

import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from supabase import create_client


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = PROJECT_ROOT / "Backend"
sys.path.insert(0, str(BACKEND_DIR))

from app.core.config import settings  # noqa: E402


SOURCE_NAME = "CONSTRUBASE_PU_48_CONSTRUCTOR"
TABLES = [
    "catalog_categories",
    "catalog_concept_aliases",
    "catalog_concept_documents",
    "catalog_concept_space_types",
    "catalog_concept_specifications",
    "catalog_concept_systems",
    "catalog_concepts",
    "catalog_construction_systems",
    "catalog_jurisdictions",
    "catalog_modules",
    "catalog_normative_sources",
    "catalog_partidas",
    "catalog_space_types",
    "catalog_system_types",
    "catalog_units",
    "engine_activation_rules",
    "engine_concept_spec_sources",
    "engine_concept_spec_space_types",
    "engine_concept_spec_systems",
    "engine_concept_specs",
    "engine_formulas",
    "engine_material_specs",
    "engine_reinforcement_specs",
    "engine_section_types",
    "engine_spec_profiles",
    "price_concept_bases",
    "price_regions",
]
MODULE_PARTIDAS = {
    "preliminares": ["PRE"],
    "cimentacion": ["CIM"],
    "estructura": ["EST"],
    "albanileria": ["ALB"],
    "instalaciones": ["HID", "SAN", "PLU", "ELE", "GAS"],
    "acabados": ["ACA"],
    "complementarios_y_equipamiento": ["CAR", "CAN", "HER", "MSA", "COM"],
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def fetch_all(client: Any, table: str, page_size: int = 1000) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    while True:
        response = client.table(table).select("*").range(offset, offset + page_size - 1).execute()
        batch = [dict(item) for item in (response.data or [])]
        rows.extend(batch)
        if len(batch) < page_size:
            return rows
        offset += page_size


def sheet_name(table: str, used: set[str]) -> str:
    candidate = table[:31]
    suffix = 1
    while candidate in used:
        tail = f"_{suffix}"
        candidate = f"{table[: 31 - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if value is None or isinstance(value, (bool, int, float, datetime)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def ordered_columns(rows: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    known: set[str] = set()
    for row in rows:
        for key in row:
            if key not in known:
                known.add(key)
                columns.append(key)
    return columns or ["sin_registros"]


def write_rows(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title)
    columns = ordered_columns(rows)
    worksheet.append(columns)
    for row in rows:
        worksheet.append([excel_value(row.get(column)) for column in columns])
    style_sheet(worksheet)


def style_sheet(worksheet: Any) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, cells in enumerate(worksheet.columns, start=1):
        max_length = max((len(str(cell.value or "")) for cell in cells), default=0)
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 10), 55)
    for cell in worksheet[1]:
        if "price" in str(cell.value).lower() or "precio" in str(cell.value).lower():
            for price_cell in worksheet[get_column_letter(cell.column)][1:]:
                if isinstance(price_cell.value, (int, float)):
                    price_cell.number_format = '$#,##0.00'


def build_used_catalog(data: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    partidas = {str(row.get("id")): row for row in data["catalog_partidas"]}
    units = {str(row.get("id")): row for row in data["catalog_units"]}
    modules_by_partida: dict[str, str] = {}
    for module, codes in MODULE_PARTIDAS.items():
        for partida in data["catalog_partidas"]:
            if str(partida.get("code") or "") in codes:
                modules_by_partida[str(partida.get("id"))] = module

    specs_by_concept: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for spec in data["engine_concept_specs"]:
        if spec.get("is_active", True):
            specs_by_concept[str(spec.get("concept_id"))].append(spec)
    for specs in specs_by_concept.values():
        specs.sort(key=lambda item: str(item.get("spec_code") or "zzzz"))

    prices_by_spec: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for price in data["price_concept_bases"]:
        if price.get("is_active", True):
            prices_by_spec[str(price.get("concept_spec_id"))].append(price)

    used: list[dict[str, Any]] = []
    for concept in data["catalog_concepts"]:
        partida_id = str(concept.get("partida_id") or "")
        module = modules_by_partida.get(partida_id)
        if not module or not concept.get("is_active", True):
            continue
        partida = partidas.get(partida_id, {})
        unit = units.get(str(concept.get("unit_id") or ""), {})
        selected_spec = (specs_by_concept.get(str(concept.get("id"))) or [{}])[0]
        prices = prices_by_spec.get(str(selected_spec.get("id"))) or []
        selected_price = next(
            (price for price in prices if str(price.get("source_name") or "") == SOURCE_NAME),
            prices[0] if prices else {},
        )
        used.append(
            {
                "modulo": module,
                "partida_codigo": partida.get("code"),
                "partida_nombre": partida.get("name"),
                "concepto_id": concept.get("id"),
                "concepto_codigo": concept.get("code"),
                "descripcion_tecnica": concept.get("technical_description"),
                "descripcion_oficial": concept.get("official_description"),
                "unidad_codigo": unit.get("code"),
                "unidad_simbolo": unit.get("symbol"),
                "formula_predeterminada": concept.get("default_formula_code"),
                "modo_cuantificacion": concept.get("quantification_mode"),
                "spec_id": selected_spec.get("id"),
                "spec_codigo": selected_spec.get("spec_code"),
                "precio_unitario": selected_price.get("unit_price", 0),
                "fuente_precio": selected_price.get("source_name"),
                "precio_actualizado": selected_price.get("updated_at"),
                "concepto_activo": concept.get("is_active"),
                "especificacion_activa": selected_spec.get("is_active"),
                "precio_activo": selected_price.get("is_active"),
            }
        )
    return sorted(used, key=lambda item: (str(item["modulo"]), str(item["concepto_codigo"])))


def main() -> None:
    if not settings.supabase_url or not settings.supabase_admin_key:
        raise RuntimeError("Faltan SUPABASE_URL o SUPABASE_SERVICE_ROLE_KEY en Backend/.env.local")

    client = create_client(settings.supabase_url, settings.supabase_admin_key)
    data = {table: fetch_all(client, table) for table in TABLES}
    catalog = build_used_catalog(data)
    missing_prices = [row for row in catalog if not float(row.get("precio_unitario") or 0)]

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Catálogo completo de conceptos Quantia"
    workbook.properties.subject = SOURCE_NAME
    workbook.properties.creator = "Quantia V2 Local"

    summary = workbook.create_sheet("Resumen")
    summary.append(["Campo", "Valor"])
    summary_rows = [
        ("Fecha de exportación", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Fuente activa", SOURCE_NAME),
        ("Proyecto Supabase", settings.supabase_url),
        ("Conceptos utilizados", len(catalog)),
        ("Conceptos sin precio", len(missing_prices)),
        ("Tablas exportadas", len(TABLES)),
    ]
    for key, value in summary_rows:
        summary.append([key, value])
    summary.append([])
    summary.append(["Tabla Supabase", "Hoja Excel", "Registros"])

    used_sheet_names = {"Resumen", "Catalogo_utilizado", "Sin_precio"}
    table_sheets: list[tuple[str, str]] = []
    for table in TABLES:
        name = sheet_name(table, used_sheet_names)
        table_sheets.append((table, name))
        summary.append([table, name, len(data[table])])
    style_sheet(summary)

    write_rows(workbook, "Catalogo_utilizado", catalog)
    write_rows(workbook, "Sin_precio", missing_prices)
    for table, name in table_sheets:
        write_rows(workbook, name, data[table])

    default_output = (
        PROJECT_ROOT.parent / "Respaldos Quantia V2" / "Catalogos"
        / f"Catalogo_Conceptos_Quantia_Completo_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )
    output = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else default_output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    verified = load_workbook(output, read_only=True, data_only=False)
    expected_sheets = 3 + len(TABLES)
    if len(verified.sheetnames) != expected_sheets:
        raise RuntimeError(f"Excel incompleto: {len(verified.sheetnames)} hojas, se esperaban {expected_sheets}")
    if verified["Catalogo_utilizado"].max_row - 1 != len(catalog):
        raise RuntimeError("La hoja consolidada no contiene todos los conceptos")
    verified.close()

    print(
        json.dumps(
            {
                "output": str(output),
                "tables": len(TABLES),
                "sheets": expected_sheets,
                "used_concepts": len(catalog),
                "missing_prices": len(missing_prices),
                "raw_rows": sum(len(rows) for rows in data.values()),
                "size_bytes": output.stat().st_size,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
